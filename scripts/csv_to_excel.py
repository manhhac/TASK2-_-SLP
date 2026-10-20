#!/usr/bin/env python
"""Convert CSV files to Excel (.xlsx)."""
import sys
from pathlib import Path
import csv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def csv_to_excel(csv_file, excel_file):
    """Convert CSV to Excel with formatting."""
    # Read CSV
    rows = []
    with open(csv_file, 'r', encoding='utf-8') as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(row)

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    # Write data
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)

            # Format header row
            if row_idx == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column].width = adjusted_width

    # Add borders
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border

    # Save
    wb.save(excel_file)
    print(f"Converted {csv_file} -> {excel_file}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python csv_to_excel.py <csv_file> [output_xlsx]")
        sys.exit(1)

    csv_file = sys.argv[1]
    excel_file = sys.argv[2] if len(sys.argv) > 2 else csv_file.replace(".csv", ".xlsx")

    csv_to_excel(csv_file, excel_file)
